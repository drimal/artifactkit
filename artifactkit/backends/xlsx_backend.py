"""Renders WorkbookSpec to .xlsx via openpyxl.

openpyxl writes formula cells as strings with no cached value, so any
formula reads back as None until something recalculates it (Excel on
open, or LibreOffice headless). validate() only checks structure here;
it does not attempt to evaluate formulas, that's a recalculation step
callers can add on top if their destination workbook must ship with
cached values (see the xlsx skill's recalc.py for the LibreOffice
approach, kept out of this backend to avoid a hard soffice dependency).

Theming: WorkbookSpec.theme, when set, fills the header row with the
accent color, sets the sheet tab color, freezes the header row, and
wraps the data range in a native Excel Table with a built-in banded
style -- reusing Excel's own named table styles rather than manually
coloring every other row, which is both less code and matches what
Excel users already recognize as "a formatted table." Only applies to
sheets that have a header; there's no sensible header row or data
range to style otherwise.

Conditional formatting: Sheet.conditional_formats renders via
openpyxl's ColorScaleRule/CellIsRule/DataBarRule -- one dispatch
method per artifactkit rule type, each translating to the matching
openpyxl rule and attached via worksheet.conditional_formatting.add().
Applied independently of theme; a sheet can have both, neither, or
just one.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import CellIsRule
from openpyxl.formatting.rule import ColorScaleRule as OpenpyxlColorScaleRule
from openpyxl.formatting.rule import DataBarRule as OpenpyxlDataBarRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from artifactkit.core.backend import ValidationResult
from artifactkit.core.models import (
    ArtifactSpec,
    CellValueRule,
    ColorScaleRule,
    DataBarRule,
    Theme,
    WorkbookSpec,
)

_HEADER_FONT = Font(bold=True)


@dataclass(frozen=True)
class _ThemeStyle:
    accent_hex: str
    font_name: str
    table_style_name: str  # one of Excel's own built-in named table styles


# Built-in Excel table style names, not artifactkit inventions -- Excel
# ships its own catalog (TableStyleLight1-21, Medium1-28, Dark1-11) and
# already knows how to render each one with banded rows and header
# formatting. Reusing them is less code than manually coloring every
# other row, and produces the exact look Excel users already recognize.
_THEME_STYLES: dict[Theme, _ThemeStyle] = {
    Theme.VIBRANT: _ThemeStyle(
        accent_hex="E94560", font_name="Calibri", table_style_name="TableStyleMedium7"
    ),
    Theme.CORPORATE: _ThemeStyle(
        accent_hex="1F3A5F", font_name="Georgia", table_style_name="TableStyleMedium2"
    ),
    Theme.MINIMAL: _ThemeStyle(
        accent_hex="999999", font_name="Helvetica", table_style_name="TableStyleLight1"
    ),
}


def _sanitize_table_name(sheet_name: str, index: int) -> str:
    """Excel table displayName must be a valid identifier (letters,
    digits, underscores; can't start with a digit) and unique across
    the workbook. Sheet names have much looser rules, so this can't
    just reuse the sheet name directly."""
    sanitized = re.sub(r"[^A-Za-z0-9_]", "_", sheet_name)
    if not sanitized or not sanitized[0].isalpha():
        sanitized = f"T{sanitized}"
    return f"Table_{index}_{sanitized}"[:255]


class XlsxBackend:
    def render(self, spec: ArtifactSpec, output_path: Path) -> None:
        assert isinstance(spec, WorkbookSpec)
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)  # default blank sheet, replaced below
        style = _THEME_STYLES.get(spec.theme) if spec.theme else None

        for sheet_index, sheet_spec in enumerate(spec.sheets, start=1):
            sheet = workbook.create_sheet(title=sheet_spec.name)
            row_offset = 0
            if sheet_spec.header:
                for col_idx, header in enumerate(sheet_spec.header, start=1):
                    cell = sheet.cell(row=1, column=col_idx, value=header)
                    cell.font = _HEADER_FONT
                row_offset = 1

            for row_idx, row in enumerate(sheet_spec.rows, start=1 + row_offset):
                for col_idx, value in enumerate(row, start=1):
                    sheet.cell(row=row_idx, column=col_idx, value=value)

            for cell_ref, formula in sheet_spec.formulas.items():
                sheet[cell_ref] = formula

            self._autosize_columns(sheet)

            if style is not None and sheet_spec.header:
                self._apply_theme(sheet, sheet_spec, style, sheet_index)

            for rule in sheet_spec.conditional_formats:
                self._apply_conditional_format(sheet, rule)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)

    def _apply_theme(self, sheet, sheet_spec, style: _ThemeStyle, sheet_index: int) -> None:
        num_cols = len(sheet_spec.header)
        num_rows = 1 + len(sheet_spec.rows)

        for col_idx in range(1, num_cols + 1):
            cell = sheet.cell(row=1, column=col_idx)
            cell.fill = PatternFill(
                start_color=style.accent_hex, end_color=style.accent_hex, fill_type="solid"
            )
            cell.font = Font(bold=True, color="FFFFFF", name=style.font_name)

        sheet.sheet_properties.tabColor = style.accent_hex
        sheet.freeze_panes = "A2"

        # A native Excel Table needs at least one data row below the
        # header; wrapping a header-only range isn't a meaningful table.
        if sheet_spec.rows:
            last_col_letter = get_column_letter(num_cols)
            table_ref = f"A1:{last_col_letter}{num_rows}"
            table = Table(
                displayName=_sanitize_table_name(sheet_spec.name, sheet_index), ref=table_ref
            )
            table.tableStyleInfo = TableStyleInfo(
                name=style.table_style_name,
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            sheet.add_table(table)

    def _apply_conditional_format(self, sheet, rule) -> None:
        if isinstance(rule, ColorScaleRule):
            self._apply_color_scale(sheet, rule)
        elif isinstance(rule, CellValueRule):
            self._apply_cell_value_rule(sheet, rule)
        elif isinstance(rule, DataBarRule):
            self._apply_data_bar(sheet, rule)
        else:  # pragma: no cover - guards against future rule types
            raise TypeError(
                f"XlsxBackend has no renderer for conditional format rule type {type(rule).__name__}"
            )

    def _apply_color_scale(self, sheet, rule: ColorScaleRule) -> None:
        if len(rule.colors) == 2:
            cf_rule = OpenpyxlColorScaleRule(
                start_type="min", start_color=rule.colors[0],
                end_type="max", end_color=rule.colors[1],
            )
        else:
            cf_rule = OpenpyxlColorScaleRule(
                start_type="min", start_color=rule.colors[0],
                mid_type="percentile", mid_value=50, mid_color=rule.colors[1],
                end_type="max", end_color=rule.colors[2],
            )
        sheet.conditional_formatting.add(rule.cell_range, cf_rule)

    def _apply_cell_value_rule(self, sheet, rule: CellValueRule) -> None:
        fill = PatternFill(start_color=rule.fill_hex, end_color=rule.fill_hex, fill_type="solid")
        font = Font(color=rule.font_hex, bold=rule.bold) if (rule.font_hex or rule.bold) else None
        cf_rule = CellIsRule(operator=rule.operator, formula=list(rule.values), fill=fill, font=font)
        sheet.conditional_formatting.add(rule.cell_range, cf_rule)

    def _apply_data_bar(self, sheet, rule: DataBarRule) -> None:
        cf_rule = OpenpyxlDataBarRule(start_type="min", end_type="max", color=rule.color_hex)
        sheet.conditional_formatting.add(rule.cell_range, cf_rule)

    def _autosize_columns(self, sheet) -> None:
        for col_cells in sheet.columns:
            length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
            col_letter = get_column_letter(col_cells[0].column)
            sheet.column_dimensions[col_letter].width = min(length + 2, 50)

    def validate(self, output_path: Path, spec: ArtifactSpec) -> ValidationResult:
        assert isinstance(spec, WorkbookSpec)
        errors: list[str] = []
        checks: list[str] = []
        try:
            workbook = openpyxl.load_workbook(output_path)
            checks.append("file_opens")
        except Exception as exc:
            return ValidationResult(is_valid=False, errors=(f"unreadable: {exc}",))

        expected_names = {s.name for s in spec.sheets}
        actual_names = set(workbook.sheetnames)
        if expected_names != actual_names:
            errors.append(f"expected sheets {expected_names}, found {actual_names}")
        else:
            checks.append("sheet_names")

        return ValidationResult(is_valid=not errors, checks_passed=tuple(checks), errors=tuple(errors))
