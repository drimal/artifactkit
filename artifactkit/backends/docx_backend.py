"""Renders DocumentSpec to .docx via python-docx.

Theming: DocumentSpec.theme, when set, colors the title and headings,
draws a thin accent-colored rule under the title, styles body text
font, and gives table header rows a filled accent background with
white bold text. python-docx has no high-level API for paragraph
borders or cell shading, so those two touches go through direct OOXML
element manipulation (_set_paragraph_bottom_border, _set_cell_shading)
-- a well-established technique for python-docx, not a hack; there is
simply no public property for either.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import docx
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, RGBColor

from artifactkit.core.backend import ValidationResult
from artifactkit.core.models import (
    ArtifactSpec,
    Block,
    DocumentSpec,
    Heading,
    Image,
    ListBlock,
    PageBreak,
    Paragraph,
    Table,
    TextStyle,
    Theme,
)


@dataclass(frozen=True)
class _ThemeStyle:
    heading_hex: str  # title + all heading levels
    body_hex: str
    accent_hex: str  # title underline rule + table header fill
    font_name: str


_THEME_STYLES: dict[Theme, _ThemeStyle] = {
    Theme.VIBRANT: _ThemeStyle(
        heading_hex="E94560", body_hex="16213E", accent_hex="F0A500", font_name="Calibri"
    ),
    Theme.CORPORATE: _ThemeStyle(
        heading_hex="1F3A5F", body_hex="333333", accent_hex="4A90D9", font_name="Georgia"
    ),
    Theme.MINIMAL: _ThemeStyle(
        heading_hex="1A1A1A", body_hex="444444", accent_hex="999999", font_name="Helvetica"
    ),
}


def _set_paragraph_bottom_border(paragraph, color_hex: str, size: int = 24) -> None:
    """python-docx exposes no paragraph-border property; this builds the
    <w:pBdr><w:bottom .../></w:pBdr> OOXML directly. size is in eighths
    of a point, per the OOXML spec (24 = 3pt)."""
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color_hex)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _set_cell_shading(cell, color_hex: str) -> None:
    """python-docx exposes no cell-fill property; this builds the
    <w:shd .../> OOXML directly, same reasoning as the border helper."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


class DocxBackend:
    def render(self, spec: ArtifactSpec, output_path: Path) -> None:
        assert isinstance(spec, DocumentSpec)
        document = docx.Document()
        document.core_properties.title = spec.title
        style = _THEME_STYLES.get(spec.theme) if spec.theme else None

        title_paragraph = document.add_heading(spec.title, level=0)
        if style is not None:
            self._style_heading_paragraph(title_paragraph, style)
            _set_paragraph_bottom_border(title_paragraph, style.accent_hex, size=24)

        for block in spec.blocks:
            self._render_block(document, block, style)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(output_path)

    def _style_heading_paragraph(self, paragraph, style: _ThemeStyle) -> None:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor.from_string(style.heading_hex)
            run.font.name = style.font_name

    def _render_block(self, document, block: Block, style: _ThemeStyle | None) -> None:
        if isinstance(block, Heading):
            paragraph = document.add_heading(block.text, level=block.level)
            if style is not None:
                self._style_heading_paragraph(paragraph, style)
        elif isinstance(block, Paragraph):
            p = document.add_paragraph()
            for run in block.runs:
                r = p.add_run(run.text)
                r.bold = run.style in (TextStyle.BOLD, TextStyle.BOLD_ITALIC)
                r.italic = run.style in (TextStyle.ITALIC, TextStyle.BOLD_ITALIC)
                if style is not None:
                    r.font.color.rgb = RGBColor.from_string(style.body_hex)
                    r.font.name = style.font_name
        elif isinstance(block, ListBlock):
            list_style = "List Number" if block.ordered else "List Bullet"
            for item in block.items:
                p = document.add_paragraph(item, style=list_style)
                if style is not None:
                    for r in p.runs:
                        r.font.color.rgb = RGBColor.from_string(style.body_hex)
                        r.font.name = style.font_name
        elif isinstance(block, Table):
            table = document.add_table(rows=1, cols=len(block.headers))
            table.style = "Light Grid Accent 1"
            for cell, header in zip(table.rows[0].cells, block.headers):
                cell.text = header
                if style is not None:
                    _set_cell_shading(cell, style.accent_hex)
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.font.color.rgb = RGBColor.from_string("FFFFFF")
                            r.font.bold = True
                            r.font.name = style.font_name
            for row in block.rows:
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    cell.text = value
                    if style is not None:
                        for p in cell.paragraphs:
                            for r in p.runs:
                                r.font.color.rgb = RGBColor.from_string(style.body_hex)
                                r.font.name = style.font_name
        elif isinstance(block, Image):
            kwargs = {"width": Inches(block.width_inches)} if block.width_inches else {}
            document.add_picture(block.source_path, **kwargs)
            if block.caption:
                caption = document.add_paragraph(block.caption)
                if "Caption" in {s.name for s in document.styles}:
                    caption.style = document.styles["Caption"]
        elif isinstance(block, PageBreak):
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        else:  # pragma: no cover - guards against future Block subtypes
            raise TypeError(f"DocxBackend has no renderer for block type {type(block).__name__}")

    def validate(self, output_path: Path, spec: ArtifactSpec) -> ValidationResult:
        assert isinstance(spec, DocumentSpec)
        errors: list[str] = []
        checks: list[str] = []
        try:
            document = docx.Document(output_path)
            checks.append("file_opens")
        except Exception as exc:
            return ValidationResult(is_valid=False, errors=(f"unreadable: {exc}",))

        expected_headings = sum(1 for b in spec.blocks if isinstance(b, Heading)) + 1  # +1 for title
        actual_headings = sum(
            1 for p in document.paragraphs if p.style.name.startswith("Heading") or p.style.name == "Title"
        )
        if actual_headings < expected_headings:
            errors.append(f"expected at least {expected_headings} headings, found {actual_headings}")
        else:
            checks.append("heading_count")

        expected_tables = sum(1 for b in spec.blocks if isinstance(b, Table))
        if len(document.tables) != expected_tables:
            errors.append(f"expected {expected_tables} tables, found {len(document.tables)}")
        else:
            checks.append("table_count")

        return ValidationResult(is_valid=not errors, checks_passed=tuple(checks), errors=tuple(errors))
