"""Tests for both adapters. These require the optional 'strands' and
'mcp' extras (pip install artifactkit[strands,mcp]) -- skipped cleanly
if they aren't installed, since most consumers only need one adapter
or none at all."""

import base64

import pytest

strands = pytest.importorskip("strands", reason="requires artifactkit[strands]")
pytest.importorskip("mcp.server.fastmcp", reason="requires artifactkit[mcp] (mcp<2.0)")

from artifactkit.adapters import mcp_server, strands_tools  # noqa: E402


DOCX_SPEC = {
    "title": "Adapter test",
    "blocks": [{"type": "paragraph", "text": "hello from adapter test"}],
}


class TestStrandsAdapter:
    def test_create_docx_writes_and_returns_location(self, tmp_path):
        result = strands_tools.create_docx(
            spec=DOCX_SPEC,
            destination_uri=str(tmp_path),
            filename="out.docx",
        )
        assert "error" not in result
        assert result["location"].endswith("out.docx")
        assert result["format"] == "docx"
        assert (tmp_path / "out.docx").exists()

    def test_create_docx_bad_spec_returns_error_dict_not_raise(self, tmp_path):
        result = strands_tools.create_docx(
            spec={"blocks": []},  # missing "title"
            destination_uri=str(tmp_path),
            filename="out.docx",
        )
        assert "error" in result
        assert result["error_type"] == "ArtifactError"

    def test_deliver_files_below_threshold(self, tmp_path):
        files = [
            {"filename": "a.txt", "content_base64": base64.b64encode(b"hello").decode()},
            {"filename": "b.txt", "content_base64": base64.b64encode(b"world").decode()},
        ]
        result = strands_tools.deliver_files(files=files, destination_uri=str(tmp_path))
        assert "error" not in result
        assert result["bundled"] is False
        assert result["file_count"] == 2

    def test_deliver_files_above_threshold_bundles(self, tmp_path):
        files = [
            {"filename": f"f{i}.txt", "content_base64": base64.b64encode(str(i).encode()).decode()}
            for i in range(6)
        ]
        result = strands_tools.deliver_files(
            files=files, destination_uri=str(tmp_path), bundle_filename="bundle"
        )
        assert "error" not in result
        assert result["bundled"] is True
        assert result["locations"][0].endswith("bundle.zip")

    def test_create_pptx_and_xlsx_and_pdf_all_work(self, tmp_path):
        pptx_result = strands_tools.create_pptx(
            spec={"title": "D", "slides": [{"layout": "title", "placeholders": {"title": "Hi"}}]},
            destination_uri=str(tmp_path),
            filename="d.pptx",
        )
        assert "error" not in pptx_result

        xlsx_result = strands_tools.create_xlsx(
            spec={"sheets": [{"name": "Data", "rows": [["1", "2"]]}]},
            destination_uri=str(tmp_path),
            filename="d.xlsx",
        )
        assert "error" not in xlsx_result

        pdf_result = strands_tools.create_pdf(
            spec=DOCX_SPEC,
            destination_uri=str(tmp_path),
            filename="d.pdf",
        )
        assert "error" not in pdf_result

    def test_inspect_pptx_template_reports_layouts(self, tmp_path):
        from pptx import Presentation as PptxPresentation

        template_path = tmp_path / "default.pptx"
        PptxPresentation().save(template_path)

        result = strands_tools.inspect_pptx_template(template_path=str(template_path))
        assert "error" not in result
        layout_names = {l["name"] for l in result["layouts"]}
        assert "Two Content" in layout_names


class TestMcpAdapter:
    def test_create_docx_writes_and_returns_location(self, tmp_path):
        result = mcp_server.create_docx(
            spec=DOCX_SPEC,
            destination_uri=str(tmp_path),
            filename="out.docx",
        )
        assert "error" not in result
        assert result["location"].endswith("out.docx")
        assert (tmp_path / "out.docx").exists()

    def test_deliver_files_source_path(self, tmp_path):
        src = tmp_path / "src.txt"
        src.write_bytes(b"from disk")
        result = mcp_server.deliver_files(
            files=[{"filename": "src.txt", "source_path": str(src)}],
            destination_uri=str(tmp_path / "out"),
        )
        assert "error" not in result
        assert result["file_count"] == 1

    def test_malformed_file_bundle_returns_error_not_raise(self, tmp_path):
        result = mcp_server.deliver_files(
            files=[{"filename": "x.txt"}],  # neither content_base64 nor source_path
            destination_uri=str(tmp_path),
        )
        assert "error" in result

    def test_mcp_server_registers_all_four_document_tools_plus_deliver(self):
        tool_names = {t.name for t in mcp_server.mcp._tool_manager.list_tools()}
        assert {"create_docx", "create_pptx", "create_xlsx", "create_pdf", "deliver_files"} <= tool_names

    def test_mcp_server_registers_inspect_pptx_template(self):
        tool_names = {t.name for t in mcp_server.mcp._tool_manager.list_tools()}
        assert "inspect_pptx_template" in tool_names

    def test_inspect_pptx_template_reports_layouts(self, tmp_path):
        from pptx import Presentation as PptxPresentation

        template_path = tmp_path / "default.pptx"
        PptxPresentation().save(template_path)

        result = mcp_server.inspect_pptx_template(template_path=str(template_path))
        assert "error" not in result
        layout_names = {l["name"] for l in result["layouts"]}
        assert "Two Content" in layout_names
        two_content = next(l for l in result["layouts"] if l["name"] == "Two Content")
        assert len(two_content["placeholders"]) == 6

    def test_inspect_pptx_template_returns_error_for_missing_file(self):
        result = mcp_server.inspect_pptx_template(template_path="/tmp/does_not_exist_artifactkit.pptx")
        assert "error" in result
