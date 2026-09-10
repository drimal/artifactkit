import zipfile

import docx
import openpyxl
import pytest
from pptx import Presentation
from pypdf import PdfReader

from artifactkit import (
    ArtifactFormat,
    ArtifactService,
    DocumentSpec,
    Heading,
    ListBlock,
    Paragraph,
    PresentationSpec,
    Sheet,
    Slide,
    Table,
    WorkbookSpec,
    destination_from_uri,
)
from artifactkit.core.errors import ArtifactError


@pytest.fixture
def service():
    return ArtifactService()


@pytest.fixture
def destination(tmp_path):
    return destination_from_uri(str(tmp_path))


def test_create_docx(service, destination, tmp_path):
    spec = DocumentSpec(
        title="Report",
        blocks=(
            Heading("Section", level=1),
            Paragraph.of("Body text."),
            ListBlock(items=("a", "b"), ordered=False),
            Table(headers=("X", "Y"), rows=(("1", "2"),)),
        ),
    )
    result = service.create(spec, ArtifactFormat.DOCX, destination, filename="report.docx")
    assert result.location.endswith("report.docx")
    assert "file_opens" in result.checks_passed
    doc = docx.Document(result.location)
    assert len(doc.tables) == 1


def test_create_pptx(service, destination):
    spec = PresentationSpec(
        title="Deck",
        slides=(
            Slide(layout="title", placeholders={"title": "Hello"}),
            Slide(layout="title_and_body", placeholders={"title": "Point", "body": "Detail"}),
        ),
    )
    result = service.create(spec, ArtifactFormat.PPTX, destination, filename="deck.pptx")
    presentation = Presentation(result.location)
    assert len(presentation.slides) == 2


def test_create_xlsx_with_formula(service, destination):
    spec = WorkbookSpec(
        sheets=(
            Sheet(
                name="Data",
                header=("Item", "Qty", "Price", "Total"),
                rows=(("Widget", 3, 2.5, None),),
                formulas={"D2": "=B2*C2"},
            ),
        )
    )
    result = service.create(spec, ArtifactFormat.XLSX, destination, filename="data.xlsx")
    workbook = openpyxl.load_workbook(result.location)
    assert workbook.sheetnames == ["Data"]
    assert workbook["Data"]["D2"].value == "=B2*C2"


def test_create_pdf(service, destination):
    spec = DocumentSpec(title="Doc", blocks=(Heading("A", level=1), Paragraph.of("body")))
    result = service.create(spec, ArtifactFormat.PDF, destination, filename="doc.pdf")
    reader = PdfReader(result.location)
    assert len(reader.pages) >= 1


def test_base_filename_resolves_extension_per_format(service, destination):
    spec = DocumentSpec(title="X", blocks=(Paragraph.of("x"),), base_filename="q3")
    docx_result = service.create(spec, ArtifactFormat.DOCX, destination)
    pdf_result = service.create(spec, ArtifactFormat.PDF, destination)
    assert docx_result.location.endswith("q3.docx")
    assert pdf_result.location.endswith("q3.pdf")


def test_explicit_filename_overrides_base_filename(service, destination):
    spec = DocumentSpec(title="X", blocks=(Paragraph.of("x"),), base_filename="q3")
    result = service.create(spec, ArtifactFormat.DOCX, destination, filename="custom.docx")
    assert result.location.endswith("custom.docx")


def test_missing_filename_raises_clear_error(service, destination):
    spec = DocumentSpec(title="X", blocks=(Paragraph.of("x"),))
    with pytest.raises(ArtifactError, match="no filename"):
        service.create(spec, ArtifactFormat.DOCX, destination)


def test_create_files_below_threshold_delivers_individually(service, destination):
    from artifactkit import FileBundleSpec, RawFile

    bundle = FileBundleSpec(files=(
        RawFile("a.py", content=b"print(1)"),
        RawFile("b.md", content=b"# doc"),
    ))
    result = service.create_files(bundle, destination)
    assert not result.bundled
    assert result.file_count == 2
    assert len(result.locations) == 2


def test_create_files_above_threshold_bundles_into_zip(service, destination):
    from artifactkit import FileBundleSpec, RawFile

    bundle = FileBundleSpec(
        files=tuple(RawFile(f"f{i}.txt", content=str(i).encode()) for i in range(6)),
        bundle_filename="export",
    )
    result = service.create_files(bundle, destination)
    assert result.bundled
    assert result.locations[0].endswith("export.zip")
    with zipfile.ZipFile(result.locations[0]) as zf:
        assert sorted(zf.namelist()) == sorted(f"f{i}.txt" for i in range(6))


def test_create_files_at_exact_threshold_not_bundled(service, destination):
    from artifactkit import FileBundleSpec, RawFile

    bundle = FileBundleSpec(files=tuple(RawFile(f"f{i}.txt", content=b"x") for i in range(5)))
    result = service.create_files(bundle, destination)
    assert not result.bundled


def test_create_files_source_path(service, destination, tmp_path):
    from artifactkit import FileBundleSpec, RawFile

    source = tmp_path / "src" / "data.csv"
    source.parent.mkdir()
    source.write_bytes(b"a,b\n1,2\n")

    bundle = FileBundleSpec(files=(RawFile("data.csv", source_path=str(source)),))
    result = service.create_files(bundle, destination)
    assert result.locations[0] != str(source)
    from pathlib import Path
    assert Path(result.locations[0]).read_bytes() == b"a,b\n1,2\n"


def test_create_files_missing_source_path_fails_before_any_write(service, destination, tmp_path):
    from artifactkit import FileBundleSpec, RawFile

    bundle = FileBundleSpec(files=(
        RawFile("ok.txt", content=b"fine"),
        RawFile("missing.txt", source_path=str(tmp_path / "does_not_exist.txt")),
    ))
    with pytest.raises(ArtifactError, match="does not exist"):
        service.create_files(bundle, destination)
    # nothing should have been written -- the check runs before the loop
    assert list(tmp_path.glob("*.txt")) == []


def test_file_bundle_from_directory(service, destination, tmp_path):
    from artifactkit import FileBundleSpec

    src_dir = tmp_path / "run_output"
    src_dir.mkdir()
    (src_dir / "a.txt").write_bytes(b"1")
    (src_dir / "b.txt").write_bytes(b"2")

    bundle = FileBundleSpec.from_directory(src_dir, bundle_filename="run")
    assert len(bundle.files) == 2
    result = service.create_files(bundle, destination)
    assert not result.bundled
    assert result.file_count == 2
