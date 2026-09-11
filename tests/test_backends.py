from pathlib import Path

import pytest

from artifactkit.backends.docx_backend import DocxBackend
from artifactkit.backends.pdf_backend import PdfBackend
from artifactkit.backends.pptx_backend import PptxBackend
from artifactkit.backends.xlsx_backend import XlsxBackend
from artifactkit.core.models import (
    DocumentSpec,
    Image,
    ListBlock,
    PageBreak,
    Paragraph,
    PresentationSpec,
    Slide,
    Table,
    WorkbookSpec,
    Sheet,
)

import base64

_PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
    "+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def test_docx_validate_rejects_unreadable_file(tmp_path):
    backend = DocxBackend()
    bad_path = tmp_path / "broken.docx"
    bad_path.write_bytes(b"this is not a real docx file")
    spec = DocumentSpec(title="x", blocks=(Paragraph.of("x"),))
    result = backend.validate(bad_path, spec)
    assert not result.is_valid
    assert any("unreadable" in e for e in result.errors)


def test_docx_validate_accepts_genuine_render(tmp_path):
    backend = DocxBackend()
    spec = DocumentSpec(title="x", blocks=(Paragraph.of("hello"),))
    path = tmp_path / "ok.docx"
    backend.render(spec, path)
    result = backend.validate(path, spec)
    assert result.is_valid
    assert "file_opens" in result.checks_passed


def test_pptx_validate_rejects_unreadable_file(tmp_path):
    backend = PptxBackend()
    bad_path = tmp_path / "broken.pptx"
    bad_path.write_bytes(b"not a pptx")
    spec = PresentationSpec(title="x", slides=(Slide(layout="title", placeholders={"title": "x"}),))
    result = backend.validate(bad_path, spec)
    assert not result.is_valid


def test_pptx_validate_catches_slide_count_mismatch(tmp_path):
    backend = PptxBackend()
    spec = PresentationSpec(
        title="x",
        slides=(
            Slide(layout="title", placeholders={"title": "a"}),
            Slide(layout="title_only", placeholders={"title": "b"}),
        ),
    )
    path = tmp_path / "deck.pptx"
    backend.render(spec, path)

    # Validate against a spec that claims 3 slides when only 2 exist.
    mismatched_spec = PresentationSpec(
        title="x",
        slides=(
            Slide(layout="title", placeholders={"title": "a"}),
            Slide(layout="title_only", placeholders={"title": "b"}),
            Slide(layout="title_only", placeholders={"title": "c"}),
        ),
    )
    result = backend.validate(path, mismatched_spec)
    assert not result.is_valid
    assert any("slides" in e for e in result.errors)


def test_xlsx_validate_rejects_unreadable_file(tmp_path):
    backend = XlsxBackend()
    bad_path = tmp_path / "broken.xlsx"
    bad_path.write_bytes(b"not an xlsx")
    spec = WorkbookSpec(sheets=(Sheet(name="Data"),))
    result = backend.validate(bad_path, spec)
    assert not result.is_valid


def test_xlsx_validate_catches_sheet_name_mismatch(tmp_path):
    backend = XlsxBackend()
    spec = WorkbookSpec(sheets=(Sheet(name="Data"),))
    path = tmp_path / "wb.xlsx"
    backend.render(spec, path)

    mismatched_spec = WorkbookSpec(sheets=(Sheet(name="Different"),))
    result = backend.validate(path, mismatched_spec)
    assert not result.is_valid


def test_pdf_validate_rejects_unreadable_file(tmp_path):
    backend = PdfBackend()
    bad_path = tmp_path / "broken.pdf"
    bad_path.write_bytes(b"%PDF-not-actually-valid")
    spec = DocumentSpec(title="x", blocks=(Paragraph.of("x"),))
    result = backend.validate(bad_path, spec)
    assert not result.is_valid


def test_pdf_validate_accepts_genuine_render(tmp_path):
    backend = PdfBackend()
    spec = DocumentSpec(title="x", blocks=(Paragraph.of("hello"),))
    path = tmp_path / "ok.pdf"
    backend.render(spec, path)
    result = backend.validate(path, spec)
    assert result.is_valid
    assert "nonzero_pages" in result.checks_passed


def test_docx_render_handles_image_and_page_break(tmp_path):
    image_path = tmp_path / "img.png"
    image_path.write_bytes(_PNG_1X1)
    backend = DocxBackend()
    spec = DocumentSpec(
        title="With image",
        blocks=(
            Image(source_path=str(image_path), caption="a tiny image", width_inches=1.0),
            PageBreak(),
            Paragraph.of("after the break"),
        ),
    )
    path = tmp_path / "with_image.docx"
    backend.render(spec, path)
    result = backend.validate(path, spec)
    assert result.is_valid


def test_pdf_render_handles_all_block_types(tmp_path):
    image_path = tmp_path / "img.png"
    image_path.write_bytes(_PNG_1X1)
    backend = PdfBackend()
    spec = DocumentSpec(
        title="Everything",
        blocks=(
            ListBlock(items=("one", "two"), ordered=True),
            Table(headers=("A", "B"), rows=(("1", "2"),)),
            Image(source_path=str(image_path), caption="pic", width_inches=0.5),
            PageBreak(),
            Paragraph.of("last page"),
        ),
    )
    path = tmp_path / "everything.pdf"
    backend.render(spec, path)
    result = backend.validate(path, spec)
    assert result.is_valid
    assert len(result.checks_passed) > 0


def test_pptx_render_with_notes_and_images(tmp_path):
    image_path = tmp_path / "img.png"
    image_path.write_bytes(_PNG_1X1)
    backend = PptxBackend()
    spec = PresentationSpec(
        title="Deck",
        slides=(
            Slide(
                layout="title_only",
                placeholders={"title": "Slide with image"},
                images=(Image(source_path=str(image_path), width_inches=1.0),),
                speaker_notes="don't forget the punchline",
            ),
        ),
    )
    path = tmp_path / "deck_with_image.pptx"
    backend.render(spec, path)
    result = backend.validate(path, spec)
    assert result.is_valid


def test_pptx_theme_hero_slide_gets_gradient_and_decorative_shape(tmp_path):
    from artifactkit.core.models import Theme
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    backend = PptxBackend()
    spec = PresentationSpec(
        title="Themed",
        slides=(Slide(layout="title", placeholders={"title": "Hello"}),),
        theme=Theme.VIBRANT,
    )
    path = tmp_path / "themed.pptx"
    backend.render(spec, path)

    from pptx import Presentation
    presentation = Presentation(path)
    slide = presentation.slides[0]

    assert str(slide.background.fill.type).startswith("GRADIENT")
    stops = slide.background.fill.gradient_stops
    assert str(stops[0].color.rgb) == "E94560"
    assert str(stops[1].color.rgb) == "533483"

    title_run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
    assert str(title_run.font.color.rgb) == "FFFFFF"  # light text on gradient
    assert title_run.font.bold is True

    accent_shapes = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
    assert len(accent_shapes) == 1  # the decorative oval, no accent bar on hero slides


def test_pptx_theme_content_slide_gets_solid_background_and_accent_bar(tmp_path):
    from artifactkit.core.models import Theme
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    backend = PptxBackend()
    spec = PresentationSpec(
        title="Themed",
        slides=(
            Slide(layout="title_and_body", placeholders={"title": "Point", "body": "Detail"}),
        ),
        theme=Theme.VIBRANT,
    )
    path = tmp_path / "content.pptx"
    backend.render(spec, path)

    from pptx import Presentation
    presentation = Presentation(path)
    slide = presentation.slides[0]

    assert str(slide.background.fill.type) == "SOLID (1)"
    assert str(slide.background.fill.fore_color.rgb) == "FFFFFF"

    title_run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
    assert str(title_run.font.color.rgb) == "E94560"  # colored title on light background

    body_placeholders = [p for p in slide.placeholders if p != slide.shapes.title]
    body_run = body_placeholders[0].text_frame.paragraphs[0].runs[0]
    assert str(body_run.font.color.rgb) == "16213E"

    accent_shapes = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
    assert len(accent_shapes) == 1  # the accent bar, no decorative oval on content slides


def test_pptx_no_theme_leaves_default_styling_untouched(tmp_path):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    backend = PptxBackend()
    spec = PresentationSpec(
        title="Plain",
        slides=(Slide(layout="title", placeholders={"title": "Hello"}),),
    )
    path = tmp_path / "plain.pptx"
    backend.render(spec, path)

    from pptx import Presentation
    presentation = Presentation(path)
    slide = presentation.slides[0]
    accent_shapes = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
    assert len(accent_shapes) == 0


def test_docx_theme_applies_title_color_border_and_table_header(tmp_path):
    from artifactkit.core.models import Theme

    backend = DocxBackend()
    spec = DocumentSpec(
        title="Themed",
        blocks=(
            Paragraph.of("body text"),
            Table(headers=("A", "B"), rows=(("1", "2"),)),
        ),
        theme=Theme.VIBRANT,
    )
    path = tmp_path / "themed.docx"
    backend.render(spec, path)

    import docx as docx_module
    doc = docx_module.Document(path)

    title_run = doc.paragraphs[0].runs[0]
    assert str(title_run.font.color.rgb) == "E94560"
    assert "pBdr" in doc.paragraphs[0]._p.xml  # accent border present

    header_cell = doc.tables[0].rows[0].cells[0]
    assert "w:shd" in header_cell._tc.xml  # header shading present
    header_run = header_cell.paragraphs[0].runs[0]
    assert str(header_run.font.color.rgb) == "FFFFFF"
    assert header_run.font.bold is True


def test_docx_no_theme_leaves_default_styling_untouched(tmp_path):
    backend = DocxBackend()
    spec = DocumentSpec(title="Plain", blocks=(Paragraph.of("body"),))
    path = tmp_path / "plain.docx"
    backend.render(spec, path)

    import docx as docx_module
    doc = docx_module.Document(path)
    assert "pBdr" not in doc.paragraphs[0]._p.xml
    body_run = doc.paragraphs[1].runs[0]
    assert body_run.font.color.rgb is None  # untouched, no explicit color set


def test_xlsx_theme_applies_header_fill_table_and_tab_color(tmp_path):
    from artifactkit.core.models import Theme

    backend = XlsxBackend()
    spec = WorkbookSpec(
        sheets=(
            Sheet(name="Data", header=("A", "B"), rows=(("1", "2"), ("3", "4"))),
        ),
        theme=Theme.CORPORATE,
    )
    path = tmp_path / "themed.xlsx"
    backend.render(spec, path)

    import openpyxl as openpyxl_module
    wb = openpyxl_module.load_workbook(path)
    sheet = wb["Data"]

    assert sheet["A1"].fill.start_color.rgb.endswith("1F3A5F")
    assert sheet["A1"].font.bold is True
    assert sheet.sheet_properties.tabColor is not None
    assert sheet.freeze_panes == "A2"
    assert len(sheet.tables) == 1
    table = list(sheet.tables.values())[0]
    assert table.tableStyleInfo.name == "TableStyleMedium2"
    assert table.ref == "A1:B3"


def test_xlsx_theme_skips_headerless_sheet_without_crashing(tmp_path):
    from artifactkit.core.models import Theme

    backend = XlsxBackend()
    spec = WorkbookSpec(
        sheets=(Sheet(name="NoHeader", rows=(("a", "b"),)),),
        theme=Theme.VIBRANT,
    )
    path = tmp_path / "headerless.xlsx"
    backend.render(spec, path)  # must not raise

    import openpyxl as openpyxl_module
    wb = openpyxl_module.load_workbook(path)
    sheet = wb["NoHeader"]
    assert len(sheet.tables) == 0  # no header, no table wrapper applied


def test_xlsx_no_theme_leaves_default_styling_untouched(tmp_path):
    backend = XlsxBackend()
    spec = WorkbookSpec(sheets=(Sheet(name="Data", header=("A", "B"), rows=(("1", "2"),)),))
    path = tmp_path / "plain.xlsx"
    backend.render(spec, path)

    import openpyxl as openpyxl_module
    wb = openpyxl_module.load_workbook(path)
    sheet = wb["Data"]
    assert len(sheet.tables) == 0
    assert sheet.sheet_properties.tabColor is None


def test_inspect_template_reports_layouts_and_placeholders(tmp_path):
    from artifactkit.backends.pptx_backend import inspect_template
    from pptx import Presentation as PptxPresentation

    template_path = tmp_path / "default.pptx"
    PptxPresentation().save(template_path)

    info = inspect_template(str(template_path))
    assert len(info.layouts) == 11  # python-pptx's default template

    two_content = next(l for l in info.layouts if l.name == "Two Content")
    assert two_content.index == 3
    content_placeholders = [ph for ph in two_content.placeholders if "OBJECT" in ph.type]
    assert len(content_placeholders) == 2
    assert {ph.name for ph in content_placeholders} == {"Content Placeholder 2", "Content Placeholder 3"}


def test_pptx_precise_idx_targeting_fills_correct_placeholder(tmp_path):
    from pptx import Presentation as PptxPresentation

    template_path = tmp_path / "default.pptx"
    PptxPresentation().save(template_path)

    backend = PptxBackend()
    spec = PresentationSpec(
        title="Comparison",
        slides=(
            Slide(
                layout="Two Content",
                placeholders={"title": "T", "idx:1": "left", "idx:2": "right"},
            ),
        ),
        template_path=str(template_path),
    )
    path = tmp_path / "out.pptx"
    backend.render(spec, path)

    rendered = PptxPresentation(path)
    by_idx = {ph.placeholder_format.idx: ph for ph in rendered.slides[0].placeholders}
    assert by_idx[1].text_frame.text == "left"
    assert by_idx[2].text_frame.text == "right"


def test_pptx_precise_name_targeting_fills_correct_placeholder(tmp_path):
    from pptx import Presentation as PptxPresentation

    template_path = tmp_path / "default.pptx"
    PptxPresentation().save(template_path)

    backend = PptxBackend()
    spec = PresentationSpec(
        title="Comparison",
        slides=(
            Slide(
                layout="Two Content",
                placeholders={
                    "title": "T",
                    "Content Placeholder 2": "left by name",
                    "Content Placeholder 3": "right by name",
                },
            ),
        ),
        template_path=str(template_path),
    )
    path = tmp_path / "out.pptx"
    backend.render(spec, path)

    rendered = PptxPresentation(path)
    by_name = {ph.name: ph for ph in rendered.slides[0].placeholders}
    assert by_name["Content Placeholder 2"].text_frame.text == "left by name"
    assert by_name["Content Placeholder 3"].text_frame.text == "right by name"


def test_pptx_old_style_spec_still_works_via_positional_fallback(tmp_path):
    from pptx import Presentation as PptxPresentation

    backend = PptxBackend()
    spec = PresentationSpec(
        title="Old style",
        slides=(Slide(layout="title_and_body", placeholders={"title": "Hi", "body": "Detail"}),),
    )
    path = tmp_path / "old_style.pptx"
    backend.render(spec, path)

    rendered = PptxPresentation(path)
    slide = rendered.slides[0]
    non_title = [ph for ph in slide.placeholders if ph != slide.shapes.title]
    assert slide.shapes.title.text_frame.text == "Hi"
    assert non_title[0].text_frame.text == "Detail"


def test_pptx_theme_skipped_when_template_path_set(tmp_path, caplog):
    import logging as logging_module
    from artifactkit.core.models import Theme
    from pptx import Presentation as PptxPresentation

    caplog.set_level(logging_module.WARNING, logger="artifactkit")
    template_path = tmp_path / "default.pptx"
    PptxPresentation().save(template_path)

    backend = PptxBackend()
    spec = PresentationSpec(
        title="Conflict",
        slides=(Slide(layout="title", placeholders={"title": "Should not be gradient"}),),
        theme=Theme.VIBRANT,
        template_path=str(template_path),
    )
    path = tmp_path / "conflict.pptx"
    backend.render(spec, path)

    rendered = PptxPresentation(path)
    slide = rendered.slides[0]
    is_gradient = slide.background.fill.type is not None and str(slide.background.fill.type).startswith("GRADIENT")
    assert not is_gradient  # theme was skipped, template's own background stands

    warnings = [r for r in caplog.records if r.message == "artifact.pptx.theme_skipped_for_template"]
    assert len(warnings) == 1


def test_pptx_theme_still_applies_without_template_path(tmp_path):
    """Regression guard: the conflict-skip logic above must not
    accidentally disable theming for the normal (no template) case."""
    from artifactkit.core.models import Theme
    from pptx import Presentation as PptxPresentation

    backend = PptxBackend()
    spec = PresentationSpec(
        title="No conflict",
        slides=(Slide(layout="title", placeholders={"title": "Should be gradient"}),),
        theme=Theme.VIBRANT,
    )
    path = tmp_path / "themed.pptx"
    backend.render(spec, path)

    rendered = PptxPresentation(path)
    slide = rendered.slides[0]
    assert str(slide.background.fill.type).startswith("GRADIENT")


def test_xlsx_color_scale_rule_three_color(tmp_path):
    from artifactkit.core.models import ColorScaleRule
    import openpyxl as openpyxl_module

    backend = XlsxBackend()
    spec = WorkbookSpec(sheets=(
        Sheet(
            name="Data", header=("A", "B"), rows=(("x", 10), ("y", 90)),
            conditional_formats=(ColorScaleRule(cell_range="B2:B3", colors=("F8696B", "FFEB84", "63BE7B")),),
        ),
    ))
    path = tmp_path / "cf.xlsx"
    backend.render(spec, path)

    wb = openpyxl_module.load_workbook(path)
    ws = wb["Data"]
    all_rules = [r for rules in ws.conditional_formatting._cf_rules.values() for r in rules]
    color_scale = next(r for r in all_rules if r.type == "colorScale")
    colors = [c.rgb[-6:] for c in color_scale.colorScale.color]
    assert colors == ["F8696B", "FFEB84", "63BE7B"]


def test_xlsx_color_scale_rule_two_color(tmp_path):
    from artifactkit.core.models import ColorScaleRule
    import openpyxl as openpyxl_module

    backend = XlsxBackend()
    spec = WorkbookSpec(sheets=(
        Sheet(
            name="Data", header=("A",), rows=(("1",), ("2",)),
            conditional_formats=(ColorScaleRule(cell_range="A2:A3", colors=("FF0000", "00FF00")),),
        ),
    ))
    path = tmp_path / "cf2.xlsx"
    backend.render(spec, path)

    wb = openpyxl_module.load_workbook(path)
    ws = wb["Data"]
    all_rules = [r for rules in ws.conditional_formatting._cf_rules.values() for r in rules]
    color_scale = next(r for r in all_rules if r.type == "colorScale")
    assert len(color_scale.colorScale.cfvo) == 2


def test_xlsx_cell_value_rule(tmp_path):
    from artifactkit.core.models import CellValueRule
    import openpyxl as openpyxl_module

    backend = XlsxBackend()
    spec = WorkbookSpec(sheets=(
        Sheet(
            name="Data", header=("A", "B"), rows=(("x", 10), ("y", 90)),
            conditional_formats=(
                CellValueRule(
                    cell_range="B2:B3", operator="greaterThan", values=("50",),
                    fill_hex="C6EFCE", font_hex="006100", bold=True,
                ),
            ),
        ),
    ))
    path = tmp_path / "cf3.xlsx"
    backend.render(spec, path)

    wb = openpyxl_module.load_workbook(path)
    ws = wb["Data"]
    all_rules = [r for rules in ws.conditional_formatting._cf_rules.values() for r in rules]
    cell_is = next(r for r in all_rules if r.type == "cellIs")
    assert cell_is.operator == "greaterThan"
    assert cell_is.formula == ["50"]


def test_xlsx_data_bar_rule(tmp_path):
    from artifactkit.core.models import DataBarRule
    import openpyxl as openpyxl_module

    backend = XlsxBackend()
    spec = WorkbookSpec(sheets=(
        Sheet(
            name="Data", header=("A",), rows=(("10",), ("20",)),
            conditional_formats=(DataBarRule(cell_range="A2:A3", color_hex="638EC6"),),
        ),
    ))
    path = tmp_path / "cf4.xlsx"
    backend.render(spec, path)

    wb = openpyxl_module.load_workbook(path)
    ws = wb["Data"]
    all_rules = [r for rules in ws.conditional_formatting._cf_rules.values() for r in rules]
    data_bar = next(r for r in all_rules if r.type == "dataBar")
    assert data_bar.dataBar.color.rgb[-6:] == "638EC6"


def test_xlsx_no_conditional_formats_leaves_sheet_unaffected(tmp_path):
    import openpyxl as openpyxl_module

    backend = XlsxBackend()
    spec = WorkbookSpec(sheets=(Sheet(name="Plain", header=("A",), rows=(("1",),)),))
    path = tmp_path / "no_cf.xlsx"
    backend.render(spec, path)

    wb = openpyxl_module.load_workbook(path)
    ws = wb["Plain"]
    assert list(ws.conditional_formatting._cf_rules) == []


def test_color_scale_rule_rejects_wrong_color_count():
    from artifactkit.core.models import ColorScaleRule
    with pytest.raises(ValueError, match="2 or 3 colors"):
        ColorScaleRule(cell_range="A1:A2", colors=("FF0000",))


def test_cell_value_rule_rejects_unknown_operator():
    from artifactkit.core.models import CellValueRule
    with pytest.raises(ValueError, match="must be one of"):
        CellValueRule(cell_range="A1:A2", operator="bogus", values=("1",), fill_hex="FFFFFF")


def test_cell_value_rule_rejects_wrong_value_count_for_between():
    from artifactkit.core.models import CellValueRule
    with pytest.raises(ValueError, match="needs 2 value"):
        CellValueRule(cell_range="A1:A2", operator="between", values=("1",), fill_hex="FFFFFF")


def test_data_bar_rule_rejects_empty_range():
    from artifactkit.core.models import DataBarRule
    with pytest.raises(ValueError, match="cell_range must not be empty"):
        DataBarRule(cell_range="", color_hex="638EC6")
